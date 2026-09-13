"""Build the final Excel workbook ranking BYMA-listed companies by market cap.

Uses scripts/market_caps_raw.json (fetched from Yahoo Finance via yfinance).
Companies without a verifiable market cap are kept at the bottom, flagged as
"N/D", instead of being silently dropped or given an invented number.
"""
import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

with open("scripts/market_caps_raw.json", encoding="utf-8") as f:
    data = json.load(f)

with_cap = [r for r in data if r["market_cap"]]
without_cap = [r for r in data if not r["market_cap"]]

with_cap.sort(key=lambda r: r["market_cap"], reverse=True)
without_cap.sort(key=lambda r: r["name"])

ordered = with_cap + without_cap

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Universo BYMA"

headers = ["Rank", "Analizar (x)", "Ticker", "Ticker Yahoo Finance", "Empresa", "Sector",
           "Market Cap (ARS)", "Precio (ARS)", "Moneda"]
ws.append(headers)

header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
header_font = Font(name="Arial", color="FFFFFF", bold=True)
body_font = Font(name="Arial")
mark_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
for col in range(1, len(headers) + 1):
    c = ws.cell(row=1, column=col)
    c.fill = header_fill
    c.font = header_font
    c.alignment = Alignment(horizontal="center", vertical="center")

for i, r in enumerate(ordered, start=1):
    ws.append([
        i,
        "",
        r["ticker"],
        r["yf_ticker"],
        r["name"],
        r["sector"],
        r["market_cap"] if r["market_cap"] else "N/D",
        r["price"] if r["price"] else "N/D",
        r["currency"] if r["currency"] else "N/D",
    ])

# Column widths
widths = {"A": 6, "B": 13, "C": 10, "D": 16, "E": 42, "F": 20, "G": 18, "H": 14, "I": 10}
for col, w in widths.items():
    ws.column_dimensions[col].width = w

# Font, number format, and highlight of the fill-in column
for row in range(2, ws.max_row + 1):
    for col in range(1, len(headers) + 1):
        ws.cell(row=row, column=col).font = body_font
    ws.cell(row=row, column=2).fill = mark_fill
    cap_cell = ws.cell(row=row, column=7)
    if isinstance(cap_cell.value, (int, float)):
        cap_cell.number_format = "#,##0"
    price_cell = ws.cell(row=row, column=8)
    if isinstance(price_cell.value, (int, float)):
        price_cell.number_format = "#,##0.00"

ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"

# Notes sheet
notes = wb.create_sheet("Notas")
notes["A1"] = "Fuente y notas"
notes["A1"].font = Font(name="Arial", bold=True, size=13)
notes_text = [
    "",
    "Universo: las 82 empresas con acciones listadas en BYMA (Bolsas y Mercados Argentinos)",
    "al 13/09/2026, segun el listado oficial de emisoras:",
    "https://www.byma.com.ar/financiarse/para-emisoras/empresas-listadas/",
    "",
    "El mercado de capitales argentino no tiene 100 empresas con acciones listadas: el total",
    "oficial de BYMA es 82. Se incluyen todas para no dejar fuera ninguna, ordenadas por",
    "capitalizacion bursatil de mayor a menor.",
    "",
    "Capitalizacion de mercado (Market Cap) y precio obtenidos de Yahoo Finance (yfinance)",
    "el 13/09/2026, en pesos argentinos (ARS), salvo indicacion contraria.",
    "",
    "Empresas marcadas 'N/D': Yahoo Finance no publica capitalizacion de mercado para ellas",
    "(acciones muy poco liquidas o sin datos de acciones en circulacion). Se listan al final,",
    "en orden alfabetico, en vez de omitirlas o de asignarles un numero inventado.",
    "",
    "No se incluyen aqui companias argentinas que cotizan unicamente en el exterior",
    "(ej. MercadoLibre, Globant, Despegar), ya que no forman parte del mercado de",
    "capitales local (BYMA) sino de bolsas extranjeras.",
    "",
    "Como usar este archivo:",
    "1) En la columna 'Analizar (x)' de la hoja 'Universo BYMA', marca con una x las",
    "   empresas que quieras incluir en el analisis de cartera.",
    "2) Guarda y sube (push) el cambio a GitHub.",
    "3) Abri el notebook desde el link del README (badge 'Open in Colab') y ejecuta todas",
    "   las celdas: va a leer este mismo archivo desde GitHub y usar solo las marcadas.",
]
for i, line in enumerate(notes_text, start=2):
    c = notes.cell(row=i, column=1, value=line)
    c.font = Font(name="Arial")
notes.column_dimensions["A"].width = 100

wb.save("empresas_byma.xlsx")
print(f"Guardado empresas_byma.xlsx con {len(ordered)} empresas ({len(with_cap)} con market cap, {len(without_cap)} sin dato).")
